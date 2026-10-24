"""
generate_data.py

Generates MADE-UP sample data for "Al-Noor Foods Trading", a fictional food
wholesaler in Riyadh. No real company or customer data is used.
Running it again gives identical files (fixed random seed).

Outputs (written to ./data):
  invoices_clean.xlsx        12 months, normalized buyer names, no duplicates
  invoices_with_issues.xlsx  same data after planting data problems
  bank_statement.csv         business account movements for the same 12 months
  ground_truth.json          what was planted, used to test the checks
"""
import calendar
import json
import random
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

SEED = 42
rng = np.random.default_rng(SEED)
random.seed(SEED)

OUT = Path("data")
OUT.mkdir(exist_ok=True)

END = date(2026, 9, 30)          # last day covered by the bank statement
VAT_RATE = 0.15
PAYMENT_TERMS_DAYS = 30
OPENING_BALANCE = 2_500_000.00
MONTHS = [(2025, 10), (2025, 11), (2025, 12), (2026, 1), (2026, 2), (2026, 3),
          (2026, 4), (2026, 5), (2026, 6), (2026, 7), (2026, 8), (2026, 9)]

# Seasonality: busy before Ramadan (Jan-Feb 2026), slow in summer.
SEASON = {(2025, 10): 1.00, (2025, 11): 1.00, (2025, 12): 1.05, (2026, 1): 1.25,
          (2026, 2): 1.35, (2026, 3): 1.15, (2026, 4): 0.95, (2026, 5): 0.95,
          (2026, 6): 0.85, (2026, 7): 0.75, (2026, 8): 0.75, (2026, 9): 0.90}

DOMINANT = "Al-Raha Supermarkets"   # one chain that is ~40% of revenue
OTHER_BUYERS = [
    "Najd Fresh Market", "Sanabel Bakery", "Diwan Restaurant", "Kashta Cafe",
    "Wadi Hanifa Catering", "Oasis Mini Market", "Tuwaiq Grocers", "Al-Yasmin Baqala",
    "Falcon Burger House", "Riyadh Spice Trading", "Dar Al-Mandi", "Muntazah Cafeteria",
    "Qasr Al-Hilw Sweets", "Granada Mart", "Olaya Express Market", "Hittin Fresh Foods",
    "Rawdah Family Restaurant", "Sahari Grill", "Naseem Mini Market", "Al-Malaz Grocers",
    "Zahra Catering", "Barakah Bakery", "Corner Store 24", "Khuzama Coffee House",
    "Al-Nakheel Restaurant", "Manar Supermarket", "Bayt Al-Tamr", "Yarmouk Mini Market",
    "Salam Cafeteria", "Dunes Catering", "Hayat Grocers", "Al-Fajr Bakery",
]
SUPPLIERS = ["Gulf Staples Wholesale", "Hail Grain Mills", "Qassim Dates Packing",
             "Eastern Dairy Distribution", "Coastal Frozen Foods Supply"]

# ---------------------------------------------------------------- invoices --
def business_days(year, month):
    """Saudi work week: Sunday to Thursday (Friday and Saturday are the weekend)."""
    days = calendar.monthrange(year, month)[1]
    return [date(year, month, d) for d in range(1, days + 1)
            if date(year, month, d).weekday() not in (4, 5)]


weights = np.array([1 / (i + 1) ** 0.8 for i in range(len(OTHER_BUYERS))])
weights = weights / weights.sum()

rows = []
for (y, m) in MONTHS:
    days = business_days(y, m)
    n_other = int(rng.normal(210, 12) * SEASON[(y, m)])
    other_total = 0.0
    for _ in range(n_other):
        subtotal = float(np.clip(rng.lognormal(np.log(2600), 0.85), 150, 45_000))
        subtotal = round(subtotal, 2)
        other_total += subtotal
        rows.append({"buyer": str(rng.choice(OTHER_BUYERS, p=weights)),
                     "issue_date": random.choice(days), "subtotal": subtotal})
    # The dominant buyer: ~40% of monthly revenue, in fewer, larger invoices.
    target = other_total * (0.40 / 0.60) * float(rng.normal(1.0, 0.05))
    k = int(rng.integers(18, 26))
    for share in rng.dirichlet(np.ones(k) * 8):
        rows.append({"buyer": DOMINANT, "issue_date": random.choice(days),
                     "subtotal": round(float(target * share), 2)})

df = pd.DataFrame(rows).sort_values("issue_date", kind="stable").reset_index(drop=True)
df["invoice_no"] = [f"INV-{10001 + i}" for i in range(len(df))]
df["due_date"] = df["issue_date"].apply(lambda d: d + timedelta(days=PAYMENT_TERMS_DAYS))
df["vat"] = (df["subtotal"] * VAT_RATE).round(2)
df["total"] = (df["subtotal"] + df["vat"]).round(2)

# ---------------------------------------------------------------- payments --
def to_banking_day(d):
    """Banks do not post on Friday/Saturday; move to Sunday."""
    while d.weekday() in (4, 5):
        d += timedelta(days=1)
    return d


# 6 large invoices that never show up as a payment (unpaid, or not genuine).
big = df[(df.buyer != DOMINANT) & (df.issue_date >= date(2025, 12, 1))
         & (df.issue_date <= date(2026, 6, 30))]
big = big[big.subtotal >= big.subtotal.quantile(0.75)]
unmatched = []
for buyer in rng.permutation(big.buyer.unique()):
    unmatched.append(big[big.buyer == buyer].sample(1, random_state=SEED).invoice_no.iloc[0])
    if len(unmatched) == 6:
        break

credits = []
for r in df.itertuples():
    if r.invoice_no in unmatched:
        continue
    delay = rng.normal(12, 6) if r.buyer == DOMINANT else rng.normal(3, 8)  # days after due date
    pay = r.due_date + timedelta(days=int(round(float(np.clip(delay, -7, 40)))))
    pay = to_banking_day(max(pay, r.issue_date + timedelta(days=3)))
    if pay <= END:                                   # later payments have not happened yet
        credits.append({"date": pay, "description": f"TRANSFER FROM {r.buyer.upper()}",
                        "amount": r.total, "direction": "credit"})

debits = []
revenue_by_month = {ym: float(df[df.issue_date.map(lambda d: (d.year, d.month) == ym)].subtotal.sum())
                    for ym in MONTHS}
PRIOR_MONTH_REVENUE = 1_250_000.0          # September 2025, before the period starts
SUPPLIER_COST_RATIO = 0.80                 # purchases = 80% of sales (20% gross margin), paid next month
for i, (y, m) in enumerate(MONTHS):
    debits.append({"date": to_banking_day(date(y, m, 1)), "description": "WAREHOUSE RENT",
                   "amount": 35_000.00, "direction": "debit"})
    debits.append({"date": to_banking_day(date(y, m, 10)), "description": "UTILITIES",
                   "amount": round(float(rng.normal(7800, 900)), 2), "direction": "debit"})
    debits.append({"date": to_banking_day(date(y, m, 15)), "description": "FUEL AND LOGISTICS",
                   "amount": round(float(rng.normal(55_000, 6_000)), 2), "direction": "debit"})
    debits.append({"date": to_banking_day(date(y, m, 27)), "description": "PAYROLL",
                   "amount": 85_000.00, "direction": "debit"})
    # Suppliers are paid in the month after the sales they relate to (30-day supplier credit).
    prev_revenue = revenue_by_month[MONTHS[i - 1]] if i > 0 else PRIOR_MONTH_REVENUE
    purchases_incl_vat = prev_revenue * SUPPLIER_COST_RATIO * (1 + VAT_RATE)
    for share in rng.dirichlet(np.ones(4) * 6):
        debits.append({"date": random.choice(business_days(y, m)),
                       "description": f"TRANSFER TO {random.choice(SUPPLIERS).upper()}",
                       "amount": round(float(purchases_incl_vat * share), 2), "direction": "debit"})

# Quarterly VAT payments to the tax authority (about 3% of quarterly sales = 15% of the margin).
quarter_sales = {(2025, 10): 3_700_000.0,   # July-September 2025, before the period starts
                 (2026, 1): sum(revenue_by_month[k] for k in MONTHS[0:3]),
                 (2026, 4): sum(revenue_by_month[k] for k in MONTHS[3:6]),
                 (2026, 7): sum(revenue_by_month[k] for k in MONTHS[6:9])}
for (y, m), sales in quarter_sales.items():
    debits.append({"date": to_banking_day(date(y, m, 28)), "description": "VAT PAYMENT TO TAX AUTHORITY",
                   "amount": round(sales * 0.03, 2), "direction": "debit"})

bank = pd.DataFrame(credits + debits)
bank["order"] = (bank.direction == "debit").astype(int)   # credits first on the same day
bank = bank.sort_values(["date", "order"], kind="stable").drop(columns="order").reset_index(drop=True)
signed = np.where(bank.direction == "credit", bank.amount, -bank.amount)
bank["balance"] = (OPENING_BALANCE + np.cumsum(signed)).round(2)
bank.to_csv(OUT / "bank_statement.csv", index=False)

# ----------------------------------------------------------- invoice files --
COLUMNS = {"invoice_no": "Invoice No", "issue_date": "Issue Date", "due_date": "Due Date",
           "buyer": "Customer", "subtotal": "Subtotal (SAR)", "vat": "VAT 15% (SAR)",
           "total": "Total (SAR)"}


def save_invoices(frame, path):
    out = frame[list(COLUMNS)].rename(columns=COLUMNS)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        out.to_excel(writer, index=False, sheet_name="Invoices")
        ws = writer.sheets["Invoices"]
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="0F2A4A")
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                if cell.column >= 5:
                    cell.number_format = "#,##0.00"
        for col, width in zip("ABCDEFG", (14, 13, 13, 30, 16, 16, 16)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"


save_invoices(df, OUT / "invoices_clean.xlsx")

# ----------------------------------------------- plant data problems (issues file)
bad = df.copy()

# 1) A whole month is missing (July 2026).
bad = bad[~bad.issue_date.apply(lambda d: (d.year, d.month) == (2026, 7))]

# 2) Two invoices appear twice (same number, date and amount).
dup_pool = bad[(bad.subtotal.between(3000, 8000)) & (bad.buyer != DOMINANT)]
dup_feb = dup_pool[dup_pool.issue_date.apply(lambda d: (d.year, d.month) == (2026, 2))].sample(1, random_state=1)
dup_may = dup_pool[dup_pool.issue_date.apply(lambda d: (d.year, d.month) == (2026, 5))].sample(1, random_state=1)
duplicates = list(dup_feb.invoice_no) + list(dup_may.invoice_no)
bad = pd.concat([bad, dup_feb, dup_may]).sort_values(["invoice_no"], kind="stable").reset_index(drop=True)

# 3) The same buyer written in different ways.
VARIANTS = {
    DOMINANT: ["Al Raha Supermarkets", "AL-RAHA SUPERMARKETS", "Al-Raha Supermarket", "Al-Raha Supermarkets Co. "],
    "Najd Fresh Market": ["Najd Fresh Mkt", "NAJD FRESH MARKET"],
}
used_variants = {}
for name, variants in VARIANTS.items():
    idx = bad.index[bad.buyer == name]
    pick = rng.choice(idx, size=max(2, int(len(idx) * 0.12)), replace=False)
    for i in pick:
        bad.at[i, "buyer"] = random.choice(variants)
    used_variants[name] = sorted(set(bad.loc[pick, "buyer"]))

save_invoices(bad, OUT / "invoices_with_issues.xlsx")

# ------------------------------------------------------------- ground truth --
share = df.groupby("buyer").subtotal.sum().sort_values(ascending=False) / df.subtotal.sum()
truth = {
    "seed": SEED,
    "period": ["2025-10-01", "2026-09-30"],
    "files": {"clean": "invoices_clean.xlsx", "with_issues": "invoices_with_issues.xlsx",
              "bank": "bank_statement.csv"},
    "planted_in_issues_file": {
        "missing_month": "2026-07",
        "duplicate_invoice_numbers": duplicates,
        "buyer_name_variants": used_variants,
    },
    "present_in_both_files": {
        "dominant_buyer": DOMINANT,
        "dominant_buyer_revenue_share": round(float(share.iloc[0]), 3),
        "top3_buyers_revenue_share": round(float(share.iloc[:3].sum()), 3),
        "invoices_without_bank_payment": unmatched,
    },
    "totals_clean_file": {"invoices": int(len(df)),
                          "revenue_excl_vat_sar": round(float(df.subtotal.sum()), 2)},
}
(OUT / "ground_truth.json").write_text(json.dumps(truth, indent=2))

monthly = df.groupby(df.issue_date.apply(lambda d: f"{d.year}-{d.month:02d}")).subtotal.agg(["count", "sum"])
print(monthly.round(0).to_string())
print(json.dumps(truth, indent=2))
